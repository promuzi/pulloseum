import sys, zipfile
from lxml import etree
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

sys.stdout.reconfigure(encoding="utf-8")
A, B = sys.argv[1], sys.argv[2]
SHIFT_SHEET, AT, N = "(없음)", 8, 2

z = zipfile.ZipFile(B)
for n in z.namelist():
    if n.endswith(".xml") or n.endswith(".rels"):
        etree.fromstring(z.read(n))
print("xml well-formed:", len(z.namelist()), "parts")
za = zipfile.ZipFile(A)
print("same part list:", za.namelist() == z.namelist())
for n in za.namelist():
    if za.read(n) != z.read(n):
        print("  changed part:", n)


def style(c):
    f, a, b = c.font, c.alignment, c.border
    col = f.color.rgb if f.color is not None and f.color.type == "rgb" else None
    return (f.name, f.sz, f.b, col, c.fill.fill_type, c.fill.fgColor.rgb, a.wrap_text, a.vertical, a.horizontal,
            b.top.style, b.bottom.style, c.number_format)


wa, wb = load_workbook(A), load_workbook(B)
print("sheets:", wa.sheetnames == wb.sheetnames)
for sa in wa.worksheets:
    sb = wb[sa.title]
    mp = (lambda r: r + N if r >= AT else r) if sa.title == SHIFT_SHEET else (lambda r: r)
    changed, styled = [], []
    seen = set()
    for row in sa.iter_rows():
        for c in row:
            r2 = mp(c.row)
            d = sb.cell(row=r2, column=c.column)
            seen.add((r2, c.column))
            if (c.value or None) != (d.value or None):
                changed.append(d.coordinate)
            if style(c) != style(d):
                styled.append((d.coordinate, style(c), style(d)))
    added = []
    for row in sb.iter_rows():
        for d in row:
            if (d.row, d.column) not in seen and d.value is not None:
                added.append(d.coordinate)
    ma = sorted(str(m) for m in sa.merged_cells.ranges)
    mb = sorted(str(m) for m in sb.merged_cells.ranges)
    ha = {mp(r): v.height for r, v in sa.row_dimensions.items() if v.height}
    hb = {r: v.height for r, v in sb.row_dimensions.items() if v.height}
    hdiff = {r: (ha.get(r), hb.get(r)) for r in set(ha) | set(hb) if ha.get(r) != hb.get(r)}
    wa_ = {k: v.width for k, v in sa.column_dimensions.items()}
    wb_ = {k: v.width for k, v in sb.column_dimensions.items()}
    print(f"[{sa.title}] value changes={len(changed)} added={added} style diffs={len(styled)}")
    print("   changed:", changed)
    for s in styled[:12]:
        print("   style:", s)
    print("   merges v4:", ma)
    print("   merges v5:", mb)
    print("   height diffs:", dict(sorted(hdiff.items())))
    print("   widths same:", wa_ == wb_, " pane:", sa.sheet_view.pane.topLeftCell, "->", sb.sheet_view.pane.topLeftCell,
          sa.sheet_view.pane.xSplit, sa.sheet_view.pane.ySplit, "/", sb.sheet_view.pane.xSplit, sb.sheet_view.pane.ySplit)
    dva = [(d.sqref, d.formula1) for d in sa.data_validations.dataValidation]
    dvb = [(d.sqref, d.formula1) for d in sb.data_validations.dataValidation]
    if dva or dvb:
        print("   dv:", dva, "->", dvb)
