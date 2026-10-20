import json, sys
from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string, get_column_letter

sys.stdout.reconfigure(encoding="utf-8")
draft, out = sys.argv[1], sys.argv[2]
CHANGED = {
    "효과 분류": ['D5', 'D6', 'E22', 'J22', 'E23', 'J23', 'E25', 'J25', 'E26', 'E27', 'J27', 'J29', 'E30', 'J30', 'J31', 'E33', 'J33', 'E34', 'J34', 'E36', 'J36', 'J37', 'E44', 'J44'],
    "스킬 분류": ['H10', 'H17', 'H18', 'B23', 'C23', 'D23', 'E23', 'H23', 'H24', 'D26', 'D28', 'E28', 'H28', 'H29']
                + [f"{c}{r}" for r in (8, 9) for c in "BCDEFGH"],
    "검토 대기": ['C6', 'D6', 'F6', 'G6', 'C7', 'E7', 'F7', 'G7', 'C8', 'E8', 'F8', 'G8', 'G9', 'C10', 'D10', 'F10', 'G10', 'C11', 'F11', 'C12', 'F12', 'C13', 'F13', 'G13', 'C14', 'F14', 'G14', 'B15', 'C15', 'E15', 'F15', 'G15', 'C16', 'F16', 'G16', 'C17', 'F17', 'G17', 'C19', 'F19', 'C22', 'F22', 'G22', 'C23', 'F23', 'F28'],
    "미분류·잔여": ['E10'],
}
wb = load_workbook(draft)
jobs = []
for si, ws in enumerate(wb.worksheets, start=1):
    for coord in CHANGED.get(ws.title, []):
        c = ws[coord]
        width = ws.column_dimensions[c.column_letter].width
        for m in ws.merged_cells.ranges:
            if c.coordinate in m and m.min_row == m.max_row:
                width = sum(ws.column_dimensions[get_column_letter(k)].width for k in range(m.min_col, m.max_col + 1))
        jobs.append({"sheet": si, "name": ws.title, "coord": coord, "row": c.row, "width": width,
                     "bold": bool(c.font.b), "text": str(c.value),
                     "orig": ws.row_dimensions[c.row].height})
json.dump(jobs, open(out, "w", encoding="utf-8"), ensure_ascii=False, indent=0)
print("jobs", len(jobs))
