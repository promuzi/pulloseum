import json, sys
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

sys.stdout.reconfigure(encoding="utf-8")
old_path, new_path, out = sys.argv[1], sys.argv[2], sys.argv[3]
wa, wb = load_workbook(old_path), load_workbook(new_path)
jobs = []
for si, ws in enumerate(wb.worksheets, start=1):
    wo = wa[ws.title]
    for row in ws.iter_rows():
        for c in row:
            if c.value is None or c.value == wo.cell(row=c.row, column=c.column).value:
                continue
            width = ws.column_dimensions[c.column_letter].width
            for m in ws.merged_cells.ranges:
                if c.coordinate in m and m.min_row == m.max_row:
                    width = sum(ws.column_dimensions[get_column_letter(k)].width for k in range(m.min_col, m.max_col + 1))
            jobs.append({"sheet": si, "name": ws.title, "coord": c.coordinate, "row": c.row, "width": width,
                         "bold": bool(c.font.b), "text": str(c.value), "orig": ws.row_dimensions[c.row].height})
json.dump(jobs, open(out, "w", encoding="utf-8"), ensure_ascii=False, indent=0)
print("jobs", len(jobs))
